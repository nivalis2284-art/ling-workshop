"""Format-specific extraction. Segments retain stable page/paragraph/row locators."""

import csv
import json
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import OpcError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image
from pypdf import PdfReader
from pypdf.errors import PyPdfError
from pypdf.generic import ArrayObject, NullObject, StreamObject

MAX_TEXT_CHARS = 250_000
MAX_SEGMENTS = 3000
MAX_CSV_FIELD_CHARS = MAX_TEXT_CHARS
MAX_PDF_PAGE_CONTENT_BYTES = 8 * 1024 * 1024
MAX_PDF_CONTENT_BYTES = 32 * 1024 * 1024

# csv's limit is process-wide; configure it once, not around concurrent requests.
csv.field_size_limit(MAX_CSV_FIELD_CHARS)


def _inspect_pdf_page(page) -> tuple[bool, int]:
    """Inspect image resources and bound decoded page/form content before parsing text.

    get_data() still decompresses first: this is a text-processing budget, not
    a sandbox or a bound on every allocation made by the PDF decoder.
    """
    pending = [page]
    seen = set()
    has_images = False
    size = 0
    while pending:
        obj = pending.pop().get_object()
        if id(obj) in seen:
            continue
        seen.add(id(obj))
        if len(seen) > MAX_SEGMENTS:
            raise ValueError("PDF 资源超过处理限制")
        if obj.get("/Subtype") == "/Image":
            has_images = True
            continue  # Detect images without decoding their pixels.
        contents = obj if obj.get("/Subtype") == "/Form" else obj.get("/Contents")
        if contents is not None:
            contents = contents.get_object()
            streams = contents if isinstance(contents, ArrayObject) else [contents]
            for stream in streams:
                stream = stream.get_object()
                if isinstance(stream, NullObject):
                    continue
                if not isinstance(stream, StreamObject):
                    raise TypeError("PDF 内容流无效")
                size += len(stream.get_data())
                if size > MAX_PDF_PAGE_CONTENT_BYTES:
                    raise ValueError("PDF 内容超过处理限制")
        resources = obj.get("/Resources")
        if resources is not None and not isinstance(resources, NullObject):
            xobjects = resources.get_object().get("/XObject")
            if xobjects is not None and not isinstance(xobjects, NullObject):
                pending.extend(xobjects.get_object().values())
    return has_images, size


def extract(extension: str, data: bytes) -> dict:
    result = {"status": "ready", "segments": [], "metadata": {}, "warnings": []}
    total = 0
    truncated = False

    def add(location: str, text: str) -> bool:
        nonlocal total, truncated
        if not text.strip():
            return True
        if len(result["segments"]) >= MAX_SEGMENTS or total + len(text) > MAX_TEXT_CHARS:
            truncated = True
            return False
        result["segments"].append({"location": location, "text": text})
        total += len(text)
        return True

    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            with Image.open(BytesIO(data)) as image:
                width, height = image.size
                expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[
                    extension
                ]
                if image.format != expected or width * height > 25_000_000:
                    raise ValueError("图片格式或尺寸不支持")
                image.verify()
            result["metadata"] = {"width": width, "height": height}
            result["status"] = "needs_vision"
        except (OSError, Image.DecompressionBombError) as error:
            raise ValueError("图片无法读取") from error
    elif extension in {".txt", ".md", ".csv"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("文本需要 UTF-8 编码") from error
        rows = (
            enumerate(csv.reader(StringIO(text)), 1)
            if extension == ".csv"
            else enumerate(text.splitlines(), 1)
        )
        try:
            for i, value in rows:
                if not add(
                    f"line:{i}",
                    json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value,
                ):
                    break
        except csv.Error:
            raise ValueError("CSV 格式无效或字段超过处理限制") from None
    elif extension in {".docx", ".xlsx"}:
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("解压后的文档超过处理限制")
                if any("vbaproject" in name.lower() for name in archive.namelist()):
                    raise ValueError("不处理带宏的文档")
        except BadZipFile as error:
            raise ValueError("文档容器损坏") from error
        try:
            if extension == ".docx":
                document = Document(BytesIO(data))
                for i, paragraph in enumerate(document.paragraphs, 1):
                    if not add(f"paragraph:{i}", paragraph.text):
                        break
                for t, table in enumerate(document.tables, 1):
                    for r, row in enumerate(table.rows, 1):
                        if not add(
                            f"table:{t}:row:{r}", " | ".join(cell.text for cell in row.cells)
                        ):
                            break
            else:
                workbook = load_workbook(
                    BytesIO(data), read_only=True, data_only=True, keep_links=False
                )
                try:
                    for sheet in workbook:
                        for i, row in enumerate(
                            sheet.iter_rows(
                                max_row=min(sheet.max_row or MAX_SEGMENTS, MAX_SEGMENTS),
                                max_col=min(sheet.max_column or 200, 200),
                                values_only=True,
                            ),
                            1,
                        ):
                            # Keep interior gaps and false/zero values at their columns.
                            end = len(row)
                            while end and row[end - 1] in (None, ""):
                                end -= 1
                            if not end:
                                continue
                            if not add(
                                f"sheet:{sheet.title}!row:{i}",
                                json.dumps(row[:end], ensure_ascii=False, default=str),
                            ):
                                break
                        if (
                            sheet.max_row
                            and sheet.max_row > MAX_SEGMENTS
                            or sheet.max_column
                            and sheet.max_column > 200
                        ):
                            truncated = True
                    result["warnings"].append("表格读取缓存值；不执行宏，不重新计算公式。")
                finally:
                    workbook.close()
        except (
            BadZipFile,
            OpcError,
            InvalidFileException,
            KeyError,
            OSError,
            SyntaxError,
            ValueError,
        ):
            raise ValueError("文档格式无效或内容损坏") from None
    elif extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(data))
        except (PyPdfError, ValueError, OSError):
            raise ValueError("PDF 格式无效或内容损坏") from None
        if reader.is_encrypted:
            raise ValueError("请先解密 PDF 再导入")
        vision_pages = []
        content_bytes = 0
        try:
            result["metadata"]["pages"] = len(reader.pages)
            if len(reader.pages) > 200:
                truncated = True
            for i, page in enumerate(reader.pages[:200], 1):
                has_images, page_bytes = _inspect_pdf_page(page)
                content_bytes += page_bytes
                if content_bytes > MAX_PDF_CONTENT_BYTES:
                    raise ValueError("PDF 内容超过处理限制")

                def inspect_operand(operator, _operands, _cm, _tm):
                    nonlocal has_images
                    if operator == b"INLINE IMAGE":
                        has_images = True

                text = page.extract_text(visitor_operand_before=inspect_operand) or ""
                if has_images or not text.strip():
                    vision_pages.append(i)
                if not add(f"page:{i}", text):
                    break
        except (PyPdfError, KeyError, TypeError, ValueError, OSError):
            raise ValueError("PDF 内容无法读取或超过处理限制") from None
        if vision_pages:
            result["metadata"]["pages_needing_vision"] = vision_pages
            result["warnings"].append("部分页面含图片或未提取到文本，需要视觉模型进一步读取。")
            result["status"] = "partial" if result["segments"] else "needs_vision"
        if not result["segments"]:
            result["status"] = "needs_vision"
    else:
        raise ValueError("当前不支持该文件格式")
    if truncated:
        result["status"] = "partial"
        result["warnings"].append("超过单次提取限制，原始文件保留，当前结果不完整。")
    return result
