"""Attachment chat contracts at the mocked HTTP boundary; no live credentials."""

import asyncio
import base64
import hashlib
import json
from contextlib import asynccontextmanager
from io import BytesIO

import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pydantic import ValidationError
from pypdf import PdfWriter
from test_chat_runtime import FAMILIES, KEY, success

from lingnexus.common.errors import Conflict, ContextTooLarge, DomainError, PermissionDenied
from lingnexus.models import ModelCapabilities, ModelProfile
from lingnexus.runtime import ChatRequest, ChatRuntime


def configure(app, family, *, images=False, pdf=False, name="fixture", window=32768):
    pid = app.profiles.save(
        name,
        ModelProfile(
            api_family=family,
            model=name,
            base_url=f"https://{name}.invalid",
            capabilities=ModelCapabilities(images=images, pdf=pdf),
            max_output_tokens=128,
            context_window_tokens=window,
        ),
        api_key=KEY,
    )["id"]
    if name == "fixture":
        app.profiles.select(pid, None)
    return pid


def turn(agent_id, request_id="turn", **kwargs):
    message = kwargs.pop("message", "Read the files")
    return ChatRequest(agent_id=agent_id, message=message, request_id=request_id, **kwargs)


@asynccontextmanager
async def runtime_for(app, calls, family="openai_chat_completions", handler=None):
    def transport(req):
        calls.append(json.loads(req.content))
        if handler:
            return handler(req)
        return httpx.Response(
            200, headers={"content-type": "text/event-stream"}, content=success(family)
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(transport)) as client:
        runtime = ChatRuntime(app, client=client)
        try:
            yield runtime
        finally:
            await runtime.aclose()


async def finish(runtime, value):
    run = runtime.submit(value)
    events = [e async for e in runtime.events(run["id"])]
    assert events[-1]["event"] == "complete", events
    return runtime.get_run(run["id"])


def image(app, agent_id, name="photo.png"):
    buffer = BytesIO()
    Image.new("RGB", (5, 7), "white").save(buffer, format="PNG")
    return app.attachments.add(agent_id, name, buffer.getvalue()), buffer.getvalue()


def scan(app, agent_id):
    buffer, writer = BytesIO(), PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(buffer)
    return app.attachments.add(agent_id, "scan.pdf", buffer.getvalue()), buffer.getvalue()


def blob(app, item):
    with app.database.transaction() as db:
        name = db.execute("SELECT blob_name FROM attachments WHERE id=?", (item["id"],)).fetchone()[
            0
        ]
    return app.attachments.directory / name


def snapshot(app):
    with app.database.transaction() as db:
        tables = [r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        return {t: [tuple(r) for r in db.execute(f'SELECT * FROM "{t}"')] for t in tables}


def user_content(payload, family, index=-1):
    return payload["input" if family == "openai_responses" else "messages"][index]["content"]


@pytest.mark.parametrize("family", FAMILIES)
def test_native_image_payload_and_durable_link(app, agent_id, family):
    configure(app, family, images=True)
    item, data = image(app, agent_id)
    calls = []

    async def scenario():
        async with runtime_for(app, calls, family) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            manifest = run["context_manifest"]
            evidence = manifest["attachments"][0]
            assert evidence["id"] == item["id"]
            assert evidence["sha256"] == hashlib.sha256(data).hexdigest()
            assert evidence["source"] == run["user_source"]
            assert manifest["source_counts"]["attachments"] == 1
            assert manifest["attachment_estimated_input_tokens"] >= len(base64.b64encode(data))
            with app.database.transaction() as db:
                link = dict(db.execute("SELECT * FROM conversation_attachments").fetchone())
            assert link["message_id"] == run["user_message_id"]
            assert link["conversation_id"] == run["conversation_id"]
            assert link["attachment_id"] == item["id"]
            assert link["sha256"] == item["sha256"]

    asyncio.run(scenario())
    blocks = user_content(calls[0], family)
    encoded = base64.b64encode(data).decode()
    if family == "openai_chat_completions":
        assert blocks[-1] == {
            "type": "image_url",
            "image_url": {"url": f"data:image/png;base64,{encoded}"},
        }
    elif family == "anthropic_messages":
        assert blocks[-1] == {
            "type": "image",
            "source": {"type": "base64", "media_type": "image/png", "data": encoded},
        }
    else:
        assert blocks[0] == {"type": "input_text", "text": "Read the files"}
        assert blocks[-1] == {
            "type": "input_image",
            "image_url": f"data:image/png;base64,{encoded}",
        }
    assert item["id"] in json.dumps(blocks) and item["sha256"] in json.dumps(blocks)


@pytest.mark.parametrize("family", FAMILIES)
@pytest.mark.parametrize(
    "extension,locator",
    [(".md", "line:2"), (".docx", "paragraph:2"), (".xlsx", "sheet:Sheet!row:2")],
)
def test_text_documents_keep_source_locators(app, agent_id, family, extension, locator):
    configure(app, family)
    buffer = BytesIO()
    if extension == ".docx":
        doc = Document()
        doc.add_paragraph("first")
        doc.add_paragraph("source truth")
        doc.save(buffer)
    elif extension == ".xlsx":
        book = Workbook()
        book.active.append(["first"])
        book.active.append(["source truth"])
        book.save(buffer)
    else:
        buffer.write(b"first\nsource truth")
    item = app.attachments.add(agent_id, "source" + extension, buffer.getvalue())
    calls = []

    async def scenario():
        async with runtime_for(app, calls, family) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            assert locator in run["context_manifest"]["attachments"][0]["locations"]

    asyncio.run(scenario())
    content = json.dumps(user_content(calls[0], family), ensure_ascii=False)
    assert "source truth" in content and locator in content
    assert item["id"] in content and item["sha256"] in content


def test_partial_text_is_bounded_and_warned(app, agent_id):
    configure(app, "openai_chat_completions")
    item = app.attachments.add(agent_id, "partial.txt", b"x" * 20000 + b"\n" + b"tail\n" * 3100)
    assert item["status"] == "partial"
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            source = run["context_manifest"]["attachments"][0]
            assert source["status"] == "partial"
            assert source["warnings"] and source["text_truncated"]
            assert source["locations"] == ["line:1"]

    asyncio.run(scenario())
    text = json.dumps(user_content(calls[0], "openai_chat_completions"), ensure_ascii=False)
    assert len(text.encode()) < 16000
    assert "partial" in text and "line:1" in text


@pytest.mark.parametrize("kind", ["image", "scan"])
@pytest.mark.parametrize("family", FAMILIES)
def test_unsupported_visual_has_no_effects(app, agent_id, family, kind):
    configure(app, family)
    item, _ = image(app, agent_id) if kind == "image" else scan(app, agent_id)
    calls, before = [], snapshot(app)

    async def scenario():
        async with runtime_for(app, calls, family) as runtime:
            with pytest.raises(DomainError, match="(?i)(image|pdf|visual)"):
                runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))
            assert not runtime._tasks

    asyncio.run(scenario())
    assert calls == [] and snapshot(app) == before


def test_cross_agent_attachment_is_denied_before_effects(app, agent_id):
    configure(app, "openai_chat_completions", images=True)
    item, _ = image(app, app.create_agent("other"))
    calls, before = [], snapshot(app)

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            with pytest.raises(PermissionDenied):
                runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))

    asyncio.run(scenario())
    assert calls == [] and snapshot(app) == before


@pytest.mark.parametrize("family", ["anthropic_messages", "openai_responses"])
def test_scanned_pdf_requires_explicit_native_capability(app, agent_id, family):
    configure(app, family, pdf=True)
    item, data = scan(app, agent_id)
    calls = []

    async def scenario():
        async with runtime_for(app, calls, family) as runtime:
            await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))

    asyncio.run(scenario())
    block = user_content(calls[0], family)[-1]
    encoded = base64.b64encode(data).decode()
    if family == "anthropic_messages":
        assert block == {
            "type": "document",
            "title": "scan.pdf",
            "source": {"type": "base64", "media_type": "application/pdf", "data": encoded},
        }
    else:
        assert block == {
            "type": "input_file",
            "filename": "scan.pdf",
            "file_data": f"data:application/pdf;base64,{encoded}",
        }


def test_chat_completions_has_no_generic_pdf_even_when_pdf_enabled(app, agent_id):
    configure(app, "openai_chat_completions", pdf=True, images=True)
    item, _ = scan(app, agent_id)
    calls, before = [], snapshot(app)

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            with pytest.raises(DomainError, match="PDF"):
                runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))

    asyncio.run(scenario())
    assert calls == [] and snapshot(app) == before


def test_only_selected_history_attachments_are_read(app, agent_id):
    configure(app, "openai_chat_completions")
    used = app.attachments.add(agent_id, "used.txt", b"old branch secret")
    unused = app.attachments.add(agent_id, "unused.txt", b"unused secret")
    blob(app, unused).unlink()
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            first = await finish(runtime, turn(agent_id, attachment_ids=[used["id"]]))
            edited = app.conversations.edit_user(
                first["conversation_id"], first["user_message_id"], "replacement", request_id="edit"
            )
            blob(app, used).unlink()
            next_run = await finish(
                runtime, turn(agent_id, "next", conversation_id=first["conversation_id"])
            )
            assert next_run["context_manifest"]["attachments"] == []
            assert next_run["context_sources"][0] == edited["source"].model_dump()

    asyncio.run(scenario())
    assert "old branch secret" not in json.dumps(calls[-1])
    assert "unused secret" not in json.dumps(calls)


def test_regeneration_reuses_original_files_and_rejects_incompatible_ids(app, agent_id):
    configure(app, "openai_chat_completions", images=True)
    item, _ = image(app, agent_id)
    other = app.attachments.add(agent_id, "other.txt", b"different")
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            first = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            regen = {
                "conversation_id": first["conversation_id"],
                "regenerate_message_id": first["user_message_id"],
            }
            before = snapshot(app)
            with pytest.raises(Conflict, match="(?i)attachment"):
                runtime.submit(turn(agent_id, "bad-regen", attachment_ids=[other["id"]], **regen))
            assert snapshot(app) == before
            second = await finish(runtime, turn(agent_id, "regen", **regen))
            assert first["user_source"] == second["user_source"]
            assert first["user_message_id"] == second["user_message_id"]

    asyncio.run(scenario())
    assert user_content(calls[0], "openai_chat_completions") == user_content(
        calls[1], "openai_chat_completions"
    )


@pytest.mark.parametrize("problem", ["missing", "corrupt", "byte_count"])
def test_invalid_blob_preflight_and_replay(app, agent_id, problem):
    configure(app, "openai_chat_completions")
    item = app.attachments.add(agent_id, "notes.txt", b"original")
    calls = []
    value = turn(agent_id, attachment_ids=[item["id"]])

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            first = await finish(runtime, value)
            if problem == "missing":
                blob(app, item).unlink()
            elif problem == "corrupt":
                blob(app, item).write_bytes(b"modified")
            else:
                with app.database.transaction() as db:
                    db.execute("UPDATE attachments SET byte_count=1 WHERE id=?", (item["id"],))
            assert runtime.submit(value) == first
            before = snapshot(app)
            with pytest.raises(DomainError, match="(?i)attachment"):
                runtime.submit(turn(agent_id, "follow", conversation_id=first["conversation_id"]))
            assert snapshot(app) == before

    asyncio.run(scenario())
    assert len(calls) == 1


def test_native_attachment_budget_is_conservative_and_atomic(app, agent_id):
    configure(app, "openai_chat_completions", images=True, window=2000)
    item, _ = image(app, agent_id)
    calls, before = [], snapshot(app)

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            with pytest.raises(ContextTooLarge):
                runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))

    asyncio.run(scenario())
    assert calls == [] and snapshot(app) == before


def test_fresh_backup_rejects_unsupported_image_without_silent_loss(app, agent_id):
    main = configure(app, "openai_chat_completions", images=True)
    backup = configure(app, "anthropic_messages", images=True, name="backup")
    app.profiles.select(main, backup)
    item, _ = image(app, agent_id)
    calls = []

    def handler(req):
        saved = app.profiles.get(backup)
        model = ModelProfile(**saved["profile"])
        app.profiles.save(
            saved["name"],
            model.model_copy(update={"capabilities": ModelCapabilities()}),
            profile_id=backup,
            api_key=KEY,
        )
        return httpx.Response(503)

    async def scenario():
        async with runtime_for(app, calls, handler=handler) as runtime:
            run = runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))
            events = [e async for e in runtime.events(run["id"])]
            assert events[-1]["event"] == "error"
            assert not any(e["event"] == "model_switch" for e in events)
            assert runtime.get_run(run["id"])["assistant_source"] is None

    asyncio.run(scenario())
    assert len(calls) == 1
    with app.database.transaction() as db:
        assert not db.execute("SELECT 1 FROM chat_conversation_failover").fetchall()


@pytest.mark.parametrize("failover", [False, True])
def test_external_observation_kb_can_be_recalled_in_real_chat(app, agent_id, failover):
    main = configure(app, "openai_chat_completions")
    if failover:
        backup = configure(app, "openai_chat_completions", name="backup")
        app.profiles.select(main, backup)
    source = app.history.add(
        agent_id,
        "observatory schedule verified",
        kind="observation",
        scope="external",
        request_id="observed",
    )
    app.memory.knowledge.save(
        agent_id, "observatory schedule", [source], kind="fact", scope="external", request_id="kb"
    )
    calls = []

    def handler(req):
        if failover and len(calls) == 1:
            return httpx.Response(503)
        return httpx.Response(
            200,
            headers={"content-type": "text/event-stream"},
            content=success("openai_chat_completions"),
        )

    async def scenario():
        async with runtime_for(app, calls, handler=handler) as runtime:
            run = await finish(runtime, turn(agent_id, message="observatory schedule"))
            assert source.model_dump() in run["context_manifest"]["recall_sources"]
            assert source.model_dump() in run["context_sources"]

    asyncio.run(scenario())
    assert len(calls) == (2 if failover else 1)
    assert all("observatory schedule verified" in json.dumps(p) for p in calls)


def test_empty_attachment_ids_preserve_legacy_wire_fingerprint():
    args = {
        "agent_id": "agent",
        "message": "Hi",
        "profile_id": None,
        "request_id": "old",
        "effort": None,
    }
    expected = hashlib.sha256(json.dumps(args, separators=(",", ":")).encode()).hexdigest()
    for value in (ChatRequest(**args), ChatRequest(**args, attachment_ids=[])):
        dumped = value.model_dump_json(exclude={"conversation_id"})
        assert hashlib.sha256(dumped.encode()).hexdigest() == expected







@pytest.mark.parametrize("problem", ["deleted", "replaced"])
def test_blob_changed_during_stream_prevents_assistant_source(app, agent_id, problem):
    configure(app, "openai_chat_completions")
    item = app.attachments.add(agent_id, "source.txt", b"source truth")
    calls = []

    def handler(req):
        if problem == "deleted":
            blob(app, item).unlink()
        else:
            blob(app, item).write_bytes(b"changed")
        return httpx.Response(
            200,
            headers={"content-type": "text/event-stream"},
            content=success("openai_chat_completions"),
        )

    async def scenario():
        async with runtime_for(app, calls, handler=handler) as runtime:
            run = runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))
            events = [e async for e in runtime.events(run["id"])]
            assert events[-1]["event"] == "error"
            assert events[-1]["data"]["code"] == "context_invalidated"
            assert runtime.get_run(run["id"])["assistant_source"] is None

    asyncio.run(scenario())


def test_finite_history_omits_archived_attachment_without_reading_blob(app, agent_id):
    from lingnexus.runtime.payload import build_payload

    profile_id = configure(app, "openai_chat_completions")
    item = app.attachments.add(agent_id, "old.txt", b"old raw file")
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            first = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            cid = first["conversation_id"]
            for index in range(12):
                app.conversations.append_user(
                    cid, "old material " * 60, request_id=f"filler:{index}"
                )
            latest = app.conversations.append_user(cid, "latest", request_id="latest")
            blob(app, item).unlink()
            profile = ModelProfile(**app.profiles.get(profile_id)["profile"]).model_copy(
                update={"context_window_tokens": 2200}
            )
            refs = [node["source"] for node in app.conversations.messages(cid)]
            payload, manifest = build_payload(app, agent_id, profile, refs)
            assert payload["messages"][-1]["content"] == "latest"
            assert manifest["attachments"] == []
            assert first["user_source"] | {"reason": "archived_history"} in manifest["omitted"]
            assert latest["source"].model_dump() in manifest["included"]

    asyncio.run(scenario())


def test_default_capabilities_require_explicit_booleans():
    capabilities = ModelCapabilities()
    assert capabilities.images is False and capabilities.pdf is False
    for flags in ({"images": "true"}, {"pdf": 1}):
        with pytest.raises(ValidationError):
            ModelCapabilities(**flags)


def test_chat_truncation_reports_partial_even_when_upload_extraction_ready(app, agent_id):
    configure(app, "openai_chat_completions")
    item = app.attachments.add(agent_id, "complete.txt", ("长文本" * 9000).encode())
    assert item["status"] == "ready"
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            entry = run["context_manifest"]["attachments"][0]
            assert entry["status"] == "partial"
            assert entry["text_truncated"] and entry["warnings"]

    asyncio.run(scenario())


def test_aggregate_native_bytes_are_bounded_before_effects(app, agent_id, monkeypatch):
    from lingnexus.runtime import attachments

    configure(app, "openai_chat_completions", images=True)
    first, data = image(app, agent_id)
    second, _ = image(app, agent_id, name="second.png")
    monkeypatch.setattr(attachments, "MAX_NATIVE_BYTES", len(data) + 1)
    calls, before = [], snapshot(app)

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            with pytest.raises(ContextTooLarge):
                runtime.submit(turn(agent_id, attachment_ids=[first["id"], second["id"]]))

    asyncio.run(scenario())
    assert calls == [] and snapshot(app) == before


@pytest.mark.parametrize("family", FAMILIES)
def test_text_only_pdf_needs_no_native_pdf_flag(app, agent_id, family):
    from test_content import _text_page

    configure(app, family)
    writer, buffer = PdfWriter(), BytesIO()
    _text_page(writer, b"Document truth")
    writer.write(buffer)
    item = app.attachments.add(agent_id, "text.pdf", buffer.getvalue())
    calls = []

    async def scenario():
        async with runtime_for(app, calls, family) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            assert run["context_manifest"]["attachments"][0]["mode"] == "text"

    asyncio.run(scenario())
    text = json.dumps(user_content(calls[0], family), ensure_ascii=False)
    assert "Document truth" in text and "page:1" in text
    assert "input_file" not in text and "base64" not in text


@pytest.mark.parametrize("family", FAMILIES)
def test_pdf_with_text_and_visual_page_cannot_discard_visuals(app, agent_id, family):
    from test_content import _text_page

    configure(app, family, images=True)
    writer, buffer = PdfWriter(), BytesIO()
    _text_page(writer, b"Text page")
    writer.add_blank_page(width=100, height=100)
    writer.write(buffer)
    item = app.attachments.add(agent_id, "mixed.pdf", buffer.getvalue())
    assert item["status"] == "partial" and item["segments"]
    calls, before = [], snapshot(app)

    async def scenario():
        async with runtime_for(app, calls, family) as runtime:
            with pytest.raises(DomainError, match="PDF"):
                runtime.submit(turn(agent_id, attachment_ids=[item["id"]]))

    asyncio.run(scenario())
    assert calls == [] and snapshot(app) == before


@pytest.mark.parametrize("kind", ["image", "pdf"])
def test_supported_native_failover_rebuilds_files_for_new_family(app, agent_id, kind):
    family = "anthropic_messages" if kind == "image" else "openai_responses"
    main = configure(app, family, images=True, pdf=True)
    backup = configure(app, "anthropic_messages", images=True, pdf=True, name="backup")
    app.profiles.select(main, backup)
    item, data = image(app, agent_id) if kind == "image" else scan(app, agent_id)
    calls = []

    def handler(req):
        if len(calls) == 1:
            return httpx.Response(503)
        return httpx.Response(
            200,
            headers={"content-type": "text/event-stream"},
            content=success("anthropic_messages"),
        )

    async def scenario():
        async with runtime_for(app, calls, family, handler) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            assert run["profile_id"] == backup
            assert run["context_manifest"]["attachments"][0]["sha256"] == item["sha256"]

    asyncio.run(scenario())
    block = user_content(calls[-1], "anthropic_messages")[-1]
    assert base64.b64decode(block["source"]["data"]) == data


def test_attachment_history_is_reused_and_replay_ids_are_part_of_input(app, agent_id):
    configure(app, "openai_chat_completions")
    item = app.attachments.add(agent_id, "notes.txt", b"selected file truth")
    other = app.attachments.add(agent_id, "other.txt", b"other file")
    value = turn(agent_id, attachment_ids=[item["id"]])
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            first = await finish(runtime, value)
            with pytest.raises(Conflict):
                runtime.submit(value.model_copy(update={"attachment_ids": [other["id"]]}))
            second = await finish(
                runtime, turn(agent_id, "second", conversation_id=first["conversation_id"])
            )
            entry = second["context_manifest"]["attachments"][0]
            assert entry["source"] == first["user_source"]
            assert entry["id"] == item["id"]

    asyncio.run(scenario())
    assert "selected file truth" in json.dumps(calls[-1])
    assert "other file" not in json.dumps(calls)


def test_verified_blob_must_match_captured_digest_even_if_metadata_changes(app, agent_id):
    from lingnexus.runtime.attachments import describe, verify_blob

    item = app.attachments.add(agent_id, "truth.txt", b"truth one")
    plan = describe(app, agent_id, item["id"])
    replacement = b"truth two"
    blob(app, item).write_bytes(replacement)
    with app.database.transaction() as db:
        db.execute(
            "UPDATE attachments SET sha256=? WHERE id=?",
            (hashlib.sha256(replacement).hexdigest(), item["id"]),
        )
    with pytest.raises(DomainError, match="(?i)digest"):
        verify_blob(app, plan)


def test_attachment_budget_never_claims_exact_provider_token_counts(app, agent_id):
    from lingnexus.context.builder import ConservativeCounter

    class ExactCounterFlag(ConservativeCounter):
        exact = True

    configure(app, "openai_chat_completions", images=True)
    app.context.counter = ExactCounterFlag()
    item, _ = image(app, agent_id)
    calls = []

    async def scenario():
        async with runtime_for(app, calls) as runtime:
            run = await finish(runtime, turn(agent_id, attachment_ids=[item["id"]]))
            assert run["context_manifest"]["token_count_exact"] is False

    asyncio.run(scenario())
