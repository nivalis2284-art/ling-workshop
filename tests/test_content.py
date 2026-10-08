import json
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def test_text_attachment_preserves_bytes_and_line_locations(app, agent_id):
    data = "第一行\n第二行".encode()
    item = app.attachments.add(agent_id, "notes.md", data)
    assert app.attachments.read_bytes(item["id"]) == data
    assert item["segments"][1] == {"location": "line:2", "text": "第二行"}
    assert item["status"] == "ready"
    assert app.attachments.read(item["id"])["original_name"] == "notes.md"


def test_bad_or_pathlike_filename_does_not_escape_storage(app, agent_id):
    for name, data in [("../outside.txt", b"bad"), ("notes.exe", b"MZ"), ("bad.png", b"not image")]:
        with pytest.raises(ValueError):
            app.attachments.add(agent_id, name, data)
    assert app.attachments.list(agent_id) == []


def test_image_is_validated_and_kept_for_vision(app, agent_id):
    file = BytesIO()
    Image.new("RGB", (5, 7), "white").save(file, format="PNG")
    item = app.attachments.add(agent_id, "photo.png", file.getvalue())
    assert item["status"] == "needs_vision"
    assert item["metadata"]["width"] == 5
    assert item["metadata"]["height"] == 7
    assert item["segments"] == []


def test_word_and_spreadsheet_keep_segment_locations(app, agent_id):
    word, buffer = Document(), BytesIO()
    word.add_paragraph("需要保留的段落")
    word.save(buffer)
    item = app.attachments.add(agent_id, "note.docx", buffer.getvalue())
    assert item["segments"][0] == {"location": "paragraph:1", "text": "需要保留的段落"}
    book, buffer = Workbook(), BytesIO()
    book.active.title = "预算"
    book.active.append(["项目", "金额"])
    book.active.append(["材料", 25])
    book.save(buffer)
    sheet = app.attachments.add(agent_id, "budget.xlsx", buffer.getvalue())
    assert sheet["segments"][1]["location"] == "sheet:预算!row:2"
    assert "25" in sheet["segments"][1]["text"]


def test_image_only_pdf_is_not_falsely_marked_read(app, agent_id):
    writer, buffer = PdfWriter(), BytesIO()
    writer.add_blank_page(width=100, height=100)
    writer.write(buffer)
    item = app.attachments.add(agent_id, "scan.pdf", buffer.getvalue())
    assert item["status"] == "needs_vision"
    assert item["metadata"]["pages"] == 1


def test_attachment_size_limit_does_not_save_partial_data(app, agent_id):
    from lingnexus.content.service import MAX_FILE_BYTES

    with pytest.raises(ValueError):
        app.attachments.add(agent_id, "huge.txt", b"x" * (MAX_FILE_BYTES + 1))
    assert app.attachments.list(agent_id) == []


def test_narrow_spreadsheet_keeps_300_rows_and_next_sheet(app, agent_id):
    book, buffer = Workbook(), BytesIO()
    book.active.title = "Rows"
    for i in range(1, 301):
        book.active.append([f"value {i}"])
    book.create_sheet("Next").append(["next sheet"])
    book.save(buffer)

    item = app.attachments.add(agent_id, "narrow.xlsx", buffer.getvalue())

    assert item["status"] == "ready"
    assert len(item["segments"]) == 301
    assert item["segments"][0]["text"] == '["value 1"]'
    assert item["segments"][-2:] == [
        {"location": "sheet:Rows!row:300", "text": '["value 300"]'},
        {"location": "sheet:Next!row:1", "text": '["next sheet"]'},
    ]


def test_sparse_spreadsheet_skips_empty_rows_without_shifting_locations(app, agent_id):
    book, buffer = Workbook(), BytesIO()
    book.active.title = "Sparse"
    book.active.append([0, None, False, ""])
    # Formatting extends the reported width; it must not consume the text budget.
    book.active.cell(1, 200).number_format = "0.00"
    book.active["A300"] = "last sparse value"
    book.create_sheet("Next").append(["next sheet"])
    book.save(buffer)

    item = app.attachments.add(agent_id, "sparse.xlsx", buffer.getvalue())

    assert item["status"] == "ready"
    assert item["segments"] == [
        {"location": "sheet:Sparse!row:1", "text": "[0, null, false]"},
        {"location": "sheet:Sparse!row:300", "text": '["last sparse value"]'},
        {"location": "sheet:Next!row:1", "text": '["next sheet"]'},
    ]


def _text_page(writer, text=b"Page 1"):
    page = writer.add_blank_page(width=100, height=100)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    ),
                }
            ),
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 6 Tf 5 5 Td (" + text + b") Tj ET")
    page[NameObject("/Contents")] = content
    return page


def test_scan_pdf_with_footer_still_identifies_page_needing_vision(app, agent_id):
    scan = BytesIO()
    Image.new("RGB", (100, 100), "white").save(scan, format="PDF")
    writer, buffer = PdfWriter(), BytesIO()
    page = writer.add_page(PdfReader(scan).pages[0])
    page.merge_page(_text_page(PdfWriter()))
    writer.write(buffer)

    item = app.attachments.add(agent_id, "scan-footer.pdf", buffer.getvalue())

    assert item["status"] == "partial"
    assert item["metadata"]["pages_needing_vision"] == [1]
    assert item["segments"] == [{"location": "page:1", "text": "Page 1"}]
    assert item["warnings"]
    assert app.attachments.read_bytes(item["id"]) == buffer.getvalue()


def test_text_only_pdf_is_ready(app, agent_id):
    writer, buffer = PdfWriter(), BytesIO()
    _text_page(writer, b"Document text")
    writer.write(buffer)

    item = app.attachments.add(agent_id, "text.pdf", buffer.getvalue())

    assert item["status"] == "ready"
    assert item["segments"] == [{"location": "page:1", "text": "Document text"}]
    assert "pages_needing_vision" not in item["metadata"]


@pytest.mark.parametrize("extension", ["pdf", "docx", "xlsx"])
def test_malformed_document_is_a_safe_value_error_without_saving(app, agent_id, extension):
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("empty.txt", "")
    data = b"not a pdf" if extension == "pdf" else buffer.getvalue()

    with pytest.raises(ValueError) as caught:
        app.attachments.add(agent_id, f"broken.{extension}", data)

    assert not any(
        detail in str(caught.value)
        for detail in ("[Content_Types].xml", "empty.txt", "not a pdf", "Stream has ended")
    )
    assert app.attachments.list(agent_id) == []
    assert not app.attachments.directory.exists()


def test_csv_accepts_140k_field_within_text_budget(app, agent_id):
    field = "x" * 140_000

    item = app.attachments.add(agent_id, "large-field.csv", (field + ",tail\n").encode())

    assert item["status"] == "ready"
    assert len(item["segments"]) == 1
    assert item["segments"][0]["location"] == "line:1"
    assert json.loads(item["segments"][0]["text"]) == [field, "tail"]


def test_csv_rejects_field_over_explicit_limit_without_saving(app, agent_id):
    from lingnexus.content.extract import MAX_TEXT_CHARS

    with pytest.raises(ValueError):
        app.attachments.add(agent_id, "oversize-field.csv", b"x" * (MAX_TEXT_CHARS + 1))

    assert app.attachments.list(agent_id) == []
    assert not app.attachments.directory.exists()


def test_compressed_pdf_content_has_a_decoded_size_limit(app, agent_id):
    writer, buffer = PdfWriter(), BytesIO()
    page = _text_page(writer)
    content = DecodedStreamObject()
    content.set_data(b" " * (8 * 1024 * 1024 + 1))
    page[NameObject("/Contents")] = content.flate_encode()
    writer.write(buffer)
    assert len(buffer.getvalue()) < 25 * 1024 * 1024

    with pytest.raises(ValueError):
        app.attachments.add(agent_id, "expanded.pdf", buffer.getvalue())

    assert app.attachments.list(agent_id) == []
    assert not app.attachments.directory.exists()
